import io
import json
import os
import re
import time
from typing import List, Dict, Optional, Tuple

import numpy as np
import streamlit as st
import fitz  # PyMuPDF

from sentence_transformers import SentenceTransformer
from groq import Groq

try:
    import faiss
except ImportError:
    faiss = None

try:
    import docx  # python-docx
except ImportError:
    docx = None


# ============================================================
# CONFIGURATION
# ============================================================

APP_NAME = "BidReady AI"

GROQ_MODEL = "openai/gpt-oss-20b"

# Estimated Groq pricing for GROQ_MODEL, USD per 1M tokens.
# These are placeholders for the cost-per-query display below —
# update to match current published Groq pricing before quoting them.
GROQ_INPUT_COST_PER_M = 0.10
GROQ_OUTPUT_COST_PER_M = 0.50

# Transient-error retry settings for Groq calls (rate limits, network blips).
MAX_RETRIES = 3
RETRY_BACKOFF_SECONDS = 2

# Lightweight and strong general-purpose embedding model.
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

TOP_K = 6

CHUNK_SIZE = 900
CHUNK_OVERLAP = 150

# Fixed rubric for the Bid Readiness Report. Kept small and fixed
# (rather than LLM-invented) so the score is comparable across tenders.
REPORT_CATEGORIES = [
    "Eligibility",
    "Technical Match",
    "Documentation",
    "Experience",
    "Certifications",
]

REPORT_EVIDENCE_PER_CATEGORY = 5
REPORT_EVIDENCE_MAX_TOTAL = 25

STATUS_COLORS = {
    "On track": "#16A34A",
    "Needs attention": "#D97706",
    "Critical gap": "#DC2626",
}

RECOMMENDATION_COLORS = {
    "BID": "#16A34A",
    "BID WITH CONDITIONS": "#D97706",
    "NO-BID": "#DC2626",
    "NEEDS REVIEW": "#64748B",
}


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="BidReady AI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    .main-title {
        font-size: 3rem;
        font-weight: 800;
        color: #0F172A;
        margin-bottom: 0;
    }

    .subtitle {
        font-size: 1.15rem;
        color: #64748B;
        margin-bottom: 1.5rem;
    }

    .decision-card {
        padding: 25px;
        border-radius: 18px;
        text-align: center;
        margin: 10px 0 25px 0;
    }

    .score {
        font-size: 3.5rem;
        font-weight: 800;
    }

    .decision {
        font-size: 1.35rem;
        font-weight: 700;
    }

    .evidence-card {
        padding: 15px;
        border-radius: 12px;
        background: #F8FAFC;
        border: 1px solid #E2E8F0;
        margin-bottom: 12px;
    }

    .small-label {
        color: #64748B;
        font-size: 0.85rem;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SESSION STATE
# ============================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "documents" not in st.session_state:
    st.session_state.documents = []

if "index" not in st.session_state:
    st.session_state.index = None

if "chunks" not in st.session_state:
    st.session_state.chunks = []

if "embedding_model" not in st.session_state:
    st.session_state.embedding_model = None

if "documents_ready" not in st.session_state:
    st.session_state.documents_ready = False

if "report" not in st.session_state:
    st.session_state.report = None


# ============================================================
# LOAD MODELS
# ============================================================

@st.cache_resource(show_spinner="Loading embedding model...")
def load_embedding_model():
    return SentenceTransformer(EMBEDDING_MODEL)


def get_groq_client():
    """
    Get Groq API client from Streamlit secrets.

    Streamlit Cloud:
        GROQ_API_KEY = "..."

    Local environment:
        GROQ_API_KEY environment variable
    """

    api_key = None

    # Streamlit secrets
    try:
        api_key = st.secrets.get("GROQ_API_KEY")
    except Exception:
        pass

    # Environment variable fallback
    if not api_key:
        api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        raise ValueError(
            "GROQ_API_KEY was not found. "
            "Add it to Streamlit Secrets."
        )

    return Groq(api_key=api_key)


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_pages(
    uploaded_file,
    document_type: str,
) -> List[Dict]:

    pdf_bytes = uploaded_file.getvalue()

    try:
        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf",
        )
    except Exception as exc:
        raise ValueError(
            f"Could not open {uploaded_file.name}: {exc}"
        )

    pages = []

    for page_number, page in enumerate(
        document,
        start=1,
    ):

        text = page.get_text("text").strip()

        if not text:
            continue

        pages.append(
            {
                "document": uploaded_file.name,
                "document_type": document_type,
                "page": page_number,
                "text": text,
            }
        )

    document.close()

    return pages


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_pages(
    uploaded_file,
    document_type: str,
) -> List[Dict]:

    if docx is None:
        raise ValueError(
            "python-docx is not installed, so .docx files can't be "
            "read. Add 'python-docx' to requirements.txt."
        )

    docx_bytes = uploaded_file.getvalue()

    try:
        document = docx.Document(
            io.BytesIO(docx_bytes)
        )
    except Exception as exc:
        raise ValueError(
            f"Could not open {uploaded_file.name}: {exc}"
        )

    parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells if cell
            )

            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()

    if not text:
        return []

    # Word documents don't expose reliable page boundaries via
    # python-docx, so the whole document is treated as one logical
    # page; citations for .docx sources will all read "Page 1".
    return [
        {
            "document": uploaded_file.name,
            "document_type": document_type,
            "page": 1,
            "text": text,
        }
    ]


# ============================================================
# PLAIN TEXT EXTRACTION
# ============================================================

def extract_txt_pages(
    uploaded_file,
    document_type: str,
) -> List[Dict]:

    raw_bytes = uploaded_file.getvalue()

    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("utf-8", errors="replace")

    # Some plain-text exports use a form-feed character as a page
    # break; split on it if present, otherwise treat as one page.
    raw_pages = text.split("\x0c")

    pages = []

    for page_number, page_text in enumerate(raw_pages, start=1):

        page_text = page_text.strip()

        if page_text:
            pages.append(
                {
                    "document": uploaded_file.name,
                    "document_type": document_type,
                    "page": page_number,
                    "text": page_text,
                }
            )

    return pages


# ============================================================
# JSON EXTRACTION (company profile)
# ============================================================

def flatten_json_to_lines(
    data,
    prefix: str = "",
) -> List[str]:
    """Turn arbitrary JSON into readable 'path: value' lines so it
    reads like structured profile text rather than raw JSON syntax."""

    lines = []

    if isinstance(data, dict):

        for key, value in data.items():

            path = f"{prefix}.{key}" if prefix else str(key)

            if isinstance(value, (dict, list)):
                lines.extend(flatten_json_to_lines(value, path))
            else:
                lines.append(f"{path}: {value}")

    elif isinstance(data, list):

        for index, item in enumerate(data):

            path = f"{prefix}[{index}]"

            if isinstance(item, (dict, list)):
                lines.extend(flatten_json_to_lines(item, path))
            else:
                lines.append(f"{path}: {item}")

    else:
        lines.append(f"{prefix}: {data}")

    return lines


def extract_json_pages(
    uploaded_file,
    document_type: str,
) -> List[Dict]:

    raw_bytes = uploaded_file.getvalue()

    try:
        data = json.loads(raw_bytes.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(
            f"Could not parse {uploaded_file.name} as JSON: {exc}"
        )

    text = "\n".join(
        flatten_json_to_lines(data)
    ).strip()

    if not text:
        return []

    return [
        {
            "document": uploaded_file.name,
            "document_type": document_type,
            "page": 1,
            "text": text,
        }
    ]


# ============================================================
# DOCUMENT EXTRACTION DISPATCH
# ============================================================

EXTRACTORS_BY_EXTENSION = {
    "pdf": extract_pdf_pages,
    "docx": extract_docx_pages,
    "txt": extract_txt_pages,
    "json": extract_json_pages,
}


def extract_document_pages(
    uploaded_file,
    document_type: str,
) -> List[Dict]:

    extension = (
        uploaded_file.name.rsplit(".", 1)[-1].lower()
        if "." in uploaded_file.name
        else ""
    )

    extractor = EXTRACTORS_BY_EXTENSION.get(extension)

    if extractor is None:
        raise ValueError(
            f"Unsupported file type '.{extension}' for "
            f"{uploaded_file.name}."
        )

    return extractor(uploaded_file, document_type)


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


# ============================================================
# CHUNKING
# ============================================================

def chunk_page(
    page: Dict,
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> List[Dict]:

    text = clean_text(page["text"])

    if len(text) <= chunk_size:

        return [
            {
                **page,
                "chunk_id": (
                    f"{page['document']}"
                    f"-p{page['page']}-c1"
                ),
                "text": text,
            }
        ]

    chunks = []

    start = 0
    chunk_number = 1

    while start < len(text):

        end = min(
            start + chunk_size,
            len(text),
        )

        chunk_text = text[start:end].strip()

        if chunk_text:

            chunks.append(
                {
                    **page,
                    "chunk_id": (
                        f"{page['document']}"
                        f"-p{page['page']}"
                        f"-c{chunk_number}"
                    ),
                    "text": chunk_text,
                }
            )

        if end >= len(text):
            break

        start = max(
            end - overlap,
            start + 1,
        )

        chunk_number += 1

    return chunks


def create_chunks(
    pages: List[Dict],
) -> List[Dict]:

    all_chunks = []

    for page in pages:

        all_chunks.extend(
            chunk_page(page)
        )

    return all_chunks


# ============================================================
# BUILD FAISS INDEX
# ============================================================

def build_faiss_index(
    chunks: List[Dict],
    embedding_model,
):

    texts = [
        chunk["text"]
        for chunk in chunks
    ]

    embeddings = embedding_model.encode(
        texts,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    embeddings = embeddings.astype(
        "float32"
    )

    if faiss is None:
        return embeddings

    dimension = embeddings.shape[1]

    index = faiss.IndexFlatIP(
        dimension
    )

    index.add(embeddings)

    return index


# ============================================================
# RETRIEVAL
# ============================================================

def retrieve(
    query: str,
    index,
    chunks: List[Dict],
    embedding_model,
    top_k: int = TOP_K,
) -> List[Dict]:

    if index is None or not chunks:
        return []

    query_embedding = embedding_model.encode(
        [query],
        convert_to_numpy=True,
        normalize_embeddings=True,
    )

    query_embedding = query_embedding.astype(
        "float32"
    )

    result_count = min(top_k, len(chunks))

    if faiss is None:
        similarities = np.dot(index, query_embedding[0])
        indices = np.argsort(similarities)[::-1][:result_count]
        scores = similarities[indices]
    else:
        scores, indices = index.search(
            query_embedding,
            result_count,
        )

        scores = scores[0]
        indices = indices[0]

    results = []

    for score, index_number in zip(scores, indices):

        if index_number < 0:
            continue

        chunk = dict(
            chunks[index_number]
        )

        chunk["similarity"] = float(score)

        results.append(chunk)

    return results


# ============================================================
# FORMAT RETRIEVED EVIDENCE
# ============================================================

def format_context(
    results: List[Dict],
) -> str:

    if not results:
        return "No relevant evidence was retrieved."

    sections = []

    for number, result in enumerate(
        results,
        start=1,
    ):

        sections.append(
            f"""
[EVIDENCE {number}]
Document: {result['document']}
Document Type: {result['document_type']}
Page: {result['page']}
Similarity: {result['similarity']:.3f}

Content:
{result['text']}
"""
        )

    return "\n".join(sections)


# ============================================================
# GROQ ANALYSIS
# ============================================================

def estimate_cost(usage) -> float:
    """Rough USD cost estimate from a Groq usage object, using the
    placeholder per-token pricing in GROQ_INPUT_COST_PER_M / _OUTPUT_."""

    if usage is None:
        return 0.0

    prompt_tokens = getattr(usage, "prompt_tokens", 0) or 0
    completion_tokens = getattr(usage, "completion_tokens", 0) or 0

    return (
        prompt_tokens / 1_000_000 * GROQ_INPUT_COST_PER_M
        + completion_tokens / 1_000_000 * GROQ_OUTPUT_COST_PER_M
    )


def check_citation_grounding(
    answer: str,
    retrieved_results: List[Dict],
) -> Tuple[bool, List[int]]:
    """Best-effort check that every 'Page N' the model cites in its
    answer was actually present in the retrieved evidence. Flags likely
    fabricated citations; it cannot catch every phrasing, so absence of
    a warning is not a guarantee, only a lack of detected mismatch."""

    cited_pages = {
        int(match)
        for match in re.findall(r"[Pp]age\s+(\d+)", answer)
    }

    if not cited_pages:
        return True, []

    retrieved_pages = {
        result["page"] for result in retrieved_results
    }

    unverified = sorted(cited_pages - retrieved_pages)

    return (len(unverified) == 0), unverified


def _call_groq(
    system_prompt: str,
    user_prompt: str,
    max_tokens: int = 2500,
    json_mode: bool = False,
) -> Tuple[str, object]:
    """Shared Groq call with retry/backoff on transient errors. Returns
    (content, usage). Raises RuntimeError if every attempt fails."""

    client = get_groq_client()

    kwargs = {}
    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):

        try:

            response = client.chat.completions.create(
                model=GROQ_MODEL,
                messages=[
                    {
                        "role": "system",
                        "content": system_prompt,
                    },
                    {
                        "role": "user",
                        "content": user_prompt,
                    },
                ],
                temperature=0.1,
                max_tokens=max_tokens,
                **kwargs,
            )

            return (
                response.choices[0].message.content,
                response.usage,
            )

        except Exception as exc:

            last_error = exc

            if attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * attempt)

    raise RuntimeError(
        f"Groq request failed after {MAX_RETRIES} attempts: {last_error}"
    )


def analyze_with_groq(
    question: str,
    company_profile: str,
    retrieved_results: List[Dict],
) -> Tuple[str, object]:

    context = format_context(
        retrieved_results
    )

    system_prompt = """
You are BidReady AI, an expert tender-readiness assistant.

Your job is to help a company determine whether a tender appears
suitable for them.

You MUST follow these rules:

1. Use the retrieved evidence as your primary source.
2. Do not invent tender requirements.
3. Do not invent company capabilities.
4. If the evidence is insufficient, say NEEDS REVIEW.
5. Distinguish between:
   - Match
   - Partial Match
   - Gap
   - Needs Review
6. Pay special attention to mandatory requirements.
7. Cite evidence using document name and page number.
8. Never fabricate a page number.
9. Be conservative about bid recommendations.
10. Explain your reasoning clearly.

For suitability questions, consider:

- Eligibility
- Technical capability
- Relevant experience
- Certifications
- Registrations
- Financial requirements
- Documentation
- Geographic requirements
- Deadlines
- Mandatory conditions

If a mandatory requirement appears to be missing,
highlight it prominently.

A high number of matches does NOT automatically mean the tender
is suitable if an important mandatory requirement is missing.

Possible recommendations:

BID
BID WITH CONDITIONS
NO-BID
NEEDS REVIEW
"""

    user_prompt = f"""
COMPANY PROFILE
================

{company_profile}


USER QUESTION
================

{question}


RETRIEVED TENDER / COMPANY EVIDENCE
=====================================

{context}


TASK
================

Answer the user's question using the evidence above.

When appropriate, structure your response as:

1. Direct Answer
2. Readiness Assessment
3. Matching Requirements
4. Gaps / Risks
5. Evidence
6. Recommended Next Steps

If you cannot confidently determine suitability from the available
evidence, say NEEDS REVIEW rather than guessing.
"""

    return _call_groq(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        max_tokens=2500,
    )


# ============================================================
# BID READINESS REPORT
# ============================================================

def gather_report_evidence(
    chunks: List[Dict],
    index,
    embedding_model,
) -> List[Dict]:
    """Retrieve tender evidence for each fixed rubric category and
    merge into one deduplicated evidence set for the report prompt."""

    seen_chunk_ids = set()
    evidence = []

    for category in REPORT_CATEGORIES:

        results = retrieve(
            query=f"{category} requirements for this tender",
            index=index,
            chunks=chunks,
            embedding_model=embedding_model,
            top_k=REPORT_EVIDENCE_PER_CATEGORY,
        )

        for result in results:

            if result["document_type"] != "Tender":
                continue

            if result["chunk_id"] in seen_chunk_ids:
                continue

            seen_chunk_ids.add(result["chunk_id"])
            evidence.append(result)

    return evidence[:REPORT_EVIDENCE_MAX_TOTAL]


def parse_json_response(text: str) -> Optional[Dict]:
    """Best-effort parse of a model JSON response, tolerating stray
    markdown code fences or leading/trailing prose around the object."""

    cleaned = text.strip()

    cleaned = re.sub(
        r"^```(?:json)?\s*|\s*```$",
        "",
        cleaned,
        flags=re.MULTILINE,
    )

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        pass

    match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)

    if not match:
        return None

    try:
        return json.loads(match.group(0))
    except json.JSONDecodeError:
        return None


def generate_bid_readiness_report(
    company_profile: str,
    evidence_chunks: List[Dict],
) -> Tuple[Optional[Dict], str, object]:
    """Ask the LLM to reason about categories/matrix/risks/recommendation
    as JSON. Returns (parsed_dict_or_None, raw_text, usage)."""

    context = format_context(evidence_chunks)

    category_list = ", ".join(REPORT_CATEGORIES)

    system_prompt = f"""
You are BidReady AI, an expert tender-readiness analyst.

You MUST respond with a single valid JSON object and nothing else —
no markdown fences, no commentary before or after it.

You MUST follow these rules:

1. Use the retrieved evidence as your primary source.
2. Do not invent tender requirements.
3. Do not invent company capabilities.
4. Cite evidence using the exact document name and page number shown
   in the evidence blocks. Never fabricate a page number.
5. Score each category from 0-100 based on how well the company
   profile satisfies the requirements found in that category's
   evidence. If no evidence was retrieved for a category, score it
   conservatively and say so in the summary.
6. Be conservative: a high count of matched requirements does not
   justify BID if a mandatory requirement is missing.
7. Use status "Critical gap" for any category with a missing
   mandatory requirement, "Needs attention" for partial gaps, and
   "On track" otherwise.

Required JSON shape (exact keys):

{{
  "categories": [
    {{"name": "<one of: {category_list}>", "score": <0-100 int>,
      "status": "On track|Needs attention|Critical gap",
      "summary": "<one sentence>"}}
    ... one entry for EACH of these categories, in this order: {category_list}
  ],
  "compliance_matrix": [
    {{"requirement": "<short requirement text>",
      "category": "<one of: {category_list}>",
      "status": "Met|Partial|Gap|Needs Review",
      "evidence_document": "<document name or empty string>",
      "evidence_page": <page number int, or null>,
      "note": "<short note>"}}
    ... one row per distinct requirement found in the evidence
  ],
  "risks": {{
    "critical": ["<critical gap description>", ...],
    "warnings": ["<warning description>", ...],
    "strengths": ["<strength description>", ...]
  }},
  "recommendation": "BID|BID WITH CONDITIONS|NO-BID|NEEDS REVIEW",
  "recommendation_reasoning": "<2-3 sentences>",
  "action_plan": ["<concrete next step>", ...]
}}
"""

    user_prompt = f"""
COMPANY PROFILE
================

{company_profile}


RETRIEVED TENDER EVIDENCE
=====================================

{context}


TASK
================

Produce the Bid Readiness Report JSON described in the system prompt,
covering every category, using only the evidence above.
"""

    raw_text, usage = _call_groq(
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        max_tokens=4000,
        json_mode=True,
    )

    return parse_json_response(raw_text), raw_text, usage


def apply_score_rules(report: Dict) -> Dict:
    """Deterministic scoring/validation layer, kept separate from the
    LLM's reasoning per BidReady's code-vs-LLM split: the LLM proposes
    category scores and a recommendation, code computes the aggregate
    score and enforces a conservative override rule."""

    categories = report.get("categories") or []

    scores = [
        category.get("score", 0)
        for category in categories
        if isinstance(category.get("score"), (int, float))
    ]

    overall_score = round(sum(scores) / len(scores)) if scores else 0

    recommendation = report.get(
        "recommendation", "NEEDS REVIEW"
    )

    risks = report.get("risks") or {}
    critical_risks = risks.get("critical") or []

    override_note = None

    if recommendation == "BID" and critical_risks:
        recommendation = "BID WITH CONDITIONS"
        override_note = (
            "BidReady's rule engine downgraded this from BID to BID "
            "WITH CONDITIONS because one or more critical gaps were "
            "identified — resolve those before submitting."
        )

    report["overall_score"] = overall_score
    report["final_recommendation"] = recommendation
    report["override_note"] = override_note

    return report


def report_to_markdown(report: Dict) -> str:
    """Render the report dict as a portable markdown summary judges
    (or the SME) can download and keep."""

    lines = [
        "# BidReady AI — Bid Readiness Report",
        "",
        f"**Overall Score:** {report.get('overall_score', 0)}%  ",
        f"**Recommendation:** {report.get('final_recommendation', 'NEEDS REVIEW')}",
        "",
    ]

    if report.get("override_note"):
        lines.append(f"> ⚠️ {report['override_note']}")
        lines.append("")

    lines.append(f"_{report.get('recommendation_reasoning', '')}_")
    lines.append("")

    lines.append("## Category Scores")
    lines.append("")
    lines.append("| Category | Score | Status | Summary |")
    lines.append("|---|---|---|---|")

    for category in report.get("categories") or []:
        lines.append(
            f"| {category.get('name', '')} "
            f"| {category.get('score', '')}% "
            f"| {category.get('status', '')} "
            f"| {category.get('summary', '')} |"
        )

    lines.append("")
    lines.append("## Compliance Matrix")
    lines.append("")
    lines.append("| Requirement | Category | Status | Evidence | Note |")
    lines.append("|---|---|---|---|---|")

    for row in report.get("compliance_matrix") or []:
        evidence = row.get("evidence_document", "")
        page = row.get("evidence_page")
        evidence_str = f"{evidence} p.{page}" if evidence and page else evidence
        lines.append(
            f"| {row.get('requirement', '')} "
            f"| {row.get('category', '')} "
            f"| {row.get('status', '')} "
            f"| {evidence_str} "
            f"| {row.get('note', '')} |"
        )

    risks = report.get("risks") or {}

    def bullet_section(title: str, items: List[str]) -> None:
        lines.append("")
        lines.append(f"### {title}")
        if items:
            lines.extend(f"- {item}" for item in items)
        else:
            lines.append("- None")

    lines.append("")
    lines.append("## Risk Report")
    bullet_section("Critical Gaps", risks.get("critical") or [])
    bullet_section("Warnings", risks.get("warnings") or [])
    bullet_section("Strengths", risks.get("strengths") or [])

    lines.append("")
    lines.append("## Action Plan")
    lines.append("")

    for number, step in enumerate(report.get("action_plan") or [], start=1):
        lines.append(f"{number}. {step}")

    lines.append("")
    lines.append(
        "_BidReady AI is an AI-assisted decision-support tool. Always "
        "verify requirements against the original tender documentation "
        "before making a final bid decision._"
    )

    return "\n".join(lines)


# ============================================================
# DOCUMENT INGESTION
# ============================================================

def process_documents(
    tender_file,
    company_file,
    status=None,
):

    if status:
        status.update(label="Extracting text from documents...")

    tender_pages = extract_document_pages(
        tender_file,
        "Tender",
    )

    if not tender_pages:
        raise ValueError(
            f"No extractable text was found in '{tender_file.name}'. "
            "If it's a PDF, it may be scanned/image-only — try a "
            "text-based PDF, DOCX, or TXT file instead."
        )

    company_pages = extract_document_pages(
        company_file,
        "Company Profile",
    )

    if not company_pages:
        raise ValueError(
            f"No extractable text was found in '{company_file.name}'. "
            "If it's a PDF, it may be scanned/image-only — try a "
            "text-based PDF, DOCX, JSON, or TXT file instead."
        )

    all_pages = tender_pages + company_pages

    if status:
        status.update(label="Cleaning and chunking document text...")

    chunks = create_chunks(
        all_pages
    )

    if status:
        status.update(label="Loading embedding model...")

    embedding_model = load_embedding_model()

    if status:
        status.update(label="Generating embeddings and building FAISS index...")

    index = build_faiss_index(
        chunks,
        embedding_model,
    )

    return (
        all_pages,
        chunks,
        index,
    )


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="main-title">🎯 BidReady AI</div>',
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="subtitle">
        Know Before You Bid — RAG-powered tender intelligence
        for businesses.
    </div>
    """,
    unsafe_allow_html=True,
)

st.write(
    "Upload your company profile and tender, then ask questions "
    "about eligibility, requirements, risks, and bid suitability."
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header("⚙️ Configuration")

    st.info(
        "This application uses:\n\n"
        "• PyMuPDF / python-docx for PDF, DOCX, TXT, and JSON "
        "extraction\n"
        "• Sentence Transformers for embeddings\n"
        "• FAISS for vector retrieval\n"
        "• Groq for LLM reasoning"
    )

    st.divider()

    st.subheader("Required API Key")

    st.code(
        "GROQ_API_KEY",
        language="text",
    )

    st.caption(
        "Add your Groq API key through Streamlit "
        "Cloud Secrets."
    )

    st.divider()

    if st.session_state.documents_ready:

        st.success(
            "Knowledge base ready"
        )

        st.metric(
            "Indexed chunks",
            len(
                st.session_state.chunks
            ),
        )

    else:

        st.warning(
            "Upload both documents to build the knowledge base."
        )


# ============================================================
# DOCUMENT UPLOAD
# ============================================================

st.header("📄 1. Upload Documents")

col1, col2 = st.columns(2)

with col1:

    tender_file = st.file_uploader(
        "Tender / RFP Document",
        type=["pdf", "docx", "txt"],
        key="tender_doc",
        help=(
            "Upload the tender you want to evaluate "
            "(PDF, DOCX, or TXT)."
        ),
    )

with col2:

    company_file = st.file_uploader(
        "Company Profile Document",
        type=["pdf", "docx", "json", "txt"],
        key="company_doc",
        help=(
            "Upload your company's profile, capabilities, "
            "experience, certifications, etc. "
            "(PDF, DOCX, JSON, or TXT)."
        ),
    )


# ============================================================
# BUILD KNOWLEDGE BASE
# ============================================================

if tender_file and company_file:

    if st.button(
        "🔎 Build BidReady Knowledge Base",
        type="primary",
        use_container_width=True,
    ):

        with st.status(
            "Building knowledge base...",
            expanded=True,
        ) as status:

            try:

                pages, chunks, index = process_documents(
                    tender_file,
                    company_file,
                    status=status,
                )

                st.session_state.documents = pages
                st.session_state.chunks = chunks
                st.session_state.index = index
                st.session_state.embedding_model = (
                    load_embedding_model()
                )
                st.session_state.documents_ready = True

                # Clear previous conversation/report when
                # new documents are uploaded.
                st.session_state.messages = []
                st.session_state.report = None

                status.update(
                    label=f"Knowledge base ready — {len(chunks)} chunks indexed.",
                    state="complete",
                )

            except Exception as exc:

                status.update(
                    label="Knowledge base build failed.",
                    state="error",
                )

                st.error(
                    f"Could not build knowledge base: {exc}"
                )


# ============================================================
# DOCUMENT SUMMARY
# ============================================================

if st.session_state.documents_ready:

    st.header("📚 Knowledge Base")

    tender_pages_count = sum(
        1
        for item in st.session_state.documents
        if item["document_type"] == "Tender"
    )

    company_pages_count = sum(
        1
        for item in st.session_state.documents
        if item["document_type"] == "Company Profile"
    )

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Tender pages",
        tender_pages_count,
    )

    col2.metric(
        "Company profile pages",
        company_pages_count,
    )

    col3.metric(
        "Indexed chunks",
        len(st.session_state.chunks),
    )

    with st.expander(
        "🔍 View indexed document information"
    ):

        st.dataframe(
            [
                {
                    "Document": item["document"],
                    "Type": item["document_type"],
                    "Page": item["page"],
                }
                for item in st.session_state.documents
            ],
            use_container_width=True,
            hide_index=True,
        )

    with st.expander(
        "📇 View extracted company profile text (sent to the model)"
    ):

        company_profile_preview = "\n\n---\n\n".join(
            item["text"]
            for item in st.session_state.documents
            if item["document_type"] == "Company Profile"
        )

        st.caption(
            f"{len(company_profile_preview)} characters extracted "
            "from the company profile document. If this is empty, "
            "very short, or missing key facts (years of experience, "
            "certifications, past projects), the source file's text "
            "likely wasn't extracted properly (e.g. a scanned or "
            "logo/table/image-heavy PDF) and the report will show "
            "gaps for everything, even if those facts are present "
            "in the original file."
        )

        st.text_area(
            "Extracted text",
            value=company_profile_preview or "(no text extracted)",
            height=250,
            disabled=True,
            label_visibility="collapsed",
        )


# ============================================================
# BID READINESS REPORT
# ============================================================

if st.session_state.documents_ready:

    st.divider()

    st.header("📊 2. Bid Readiness Report")

    st.caption(
        "Generates a scored breakdown across "
        f"{', '.join(REPORT_CATEGORIES)}, a compliance matrix, "
        "a risk report, and an action plan."
    )

    if st.button(
        "🧮 Generate Bid Readiness Report",
        use_container_width=True,
    ):

        company_chunks = [
            chunk
            for chunk in st.session_state.chunks
            if chunk["document_type"] == "Company Profile"
        ]

        company_profile = "\n\n".join(
            chunk["text"] for chunk in company_chunks
        )

        if not company_profile.strip():

            st.warning(
                "No company profile text is indexed — re-upload a "
                "text-based company profile file (PDF, DOCX, JSON, "
                "or TXT) and rebuild the knowledge base before "
                "generating a report."
            )

        else:

            with st.spinner(
                "Scoring categories, building the compliance matrix, "
                "and assessing risk..."
            ):

                try:

                    evidence = gather_report_evidence(
                        chunks=st.session_state.chunks,
                        index=st.session_state.index,
                        embedding_model=st.session_state.embedding_model,
                    )

                    parsed, raw_text, usage = generate_bid_readiness_report(
                        company_profile=company_profile,
                        evidence_chunks=evidence,
                    )

                    if parsed is None:

                        st.error(
                            "The report could not be parsed as JSON. "
                            "Try generating it again."
                        )

                        with st.expander("Raw model output"):
                            st.text(raw_text)

                    else:

                        report = apply_score_rules(parsed)
                        st.session_state.report = report
                        st.session_state.report_usage = usage

                except Exception as exc:

                    st.error(f"Report generation failed: {exc}")

    report = st.session_state.report

    if report:

        overall_score = report.get("overall_score", 0)
        recommendation = report.get(
            "final_recommendation", "NEEDS REVIEW"
        )
        decision_color = RECOMMENDATION_COLORS.get(
            recommendation, "#64748B"
        )

        st.markdown(
            f"""
            <div class="decision-card" style="background: {decision_color}1A;
                 border: 1px solid {decision_color};">
                <div class="score" style="color: {decision_color};">
                    {overall_score}%
                </div>
                <div class="decision" style="color: {decision_color};">
                    {recommendation}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        if report.get("override_note"):
            st.warning(report["override_note"])

        if report.get("recommendation_reasoning"):
            st.write(report["recommendation_reasoning"])

        st.subheader("Category Scores")

        st.dataframe(
            [
                {
                    "Category": category.get("name", ""),
                    "Score": f"{category.get('score', 0)}%",
                    "Status": category.get("status", ""),
                    "Summary": category.get("summary", ""),
                }
                for category in report.get("categories") or []
            ],
            use_container_width=True,
            hide_index=True,
        )

        st.subheader("Compliance Matrix")

        status_icons = {
            "Met": "✅",
            "Partial": "🟡",
            "Gap": "❌",
            "Needs Review": "❓",
        }

        st.dataframe(
            [
                {
                    "Requirement": row.get("requirement", ""),
                    "Category": row.get("category", ""),
                    "Status": (
                        f"{status_icons.get(row.get('status', ''), '')} "
                        f"{row.get('status', '')}"
                    ),
                    "Evidence": (
                        f"{row.get('evidence_document', '')} "
                        f"p.{row.get('evidence_page')}"
                        if row.get("evidence_document")
                        and row.get("evidence_page")
                        else row.get("evidence_document", "")
                    ),
                    "Note": row.get("note", ""),
                }
                for row in report.get("compliance_matrix") or []
            ],
            use_container_width=True,
            hide_index=True,
        )

        st.subheader("Risk Report")

        risks = report.get("risks") or {}

        risk_col1, risk_col2, risk_col3 = st.columns(3)

        with risk_col1:
            st.markdown("**🔴 Critical Gaps**")
            for item in risks.get("critical") or []:
                st.write(f"- {item}")

        with risk_col2:
            st.markdown("**🟡 Warnings**")
            for item in risks.get("warnings") or []:
                st.write(f"- {item}")

        with risk_col3:
            st.markdown("**🟢 Strengths**")
            for item in risks.get("strengths") or []:
                st.write(f"- {item}")

        st.subheader("Action Plan")

        for number, step in enumerate(
            report.get("action_plan") or [], start=1
        ):
            st.write(f"{number}. {step}")

        report_usage = st.session_state.get("report_usage")

        if report_usage:
            st.caption(
                f"Report generation used "
                f"{getattr(report_usage, 'prompt_tokens', 0)} prompt + "
                f"{getattr(report_usage, 'completion_tokens', 0)} "
                f"completion tokens "
                f"(~${estimate_cost(report_usage):.5f} estimated)."
            )

        st.download_button(
            "⬇️ Download Report (Markdown)",
            data=report_to_markdown(report),
            file_name="bidready_report.md",
            mime="text/markdown",
            use_container_width=True,
        )


# ============================================================
# CHAT / QUESTIONS
# ============================================================

if st.session_state.documents_ready:

    st.divider()

    st.header("💬 3. Ask BidReady AI")

    st.caption(
        "Ask whether the tender is suitable, what requirements "
        "you meet, what gaps exist, or what you should do before bidding."
    )

    example_questions = [
        "Is this tender suitable for my company?",
        "What mandatory requirements does my company fail to meet?",
        "Which tender requirements does my company satisfy?",
        "What are the biggest risks if we bid?",
        "Do we have enough relevant experience?",
        "What certifications are required?",
        "What documents are missing?",
        "Should we BID, BID WITH CONDITIONS, or NO-BID?",
    ]

    selected_question = st.selectbox(
        "Example questions",
        ["Choose a question..."]
        + example_questions,
    )

    question = st.chat_input(
        "Ask BidReady AI a question..."
    )

    if (
        not question
        and selected_question != "Choose a question..."
    ):
        question = selected_question

    if question:

        # Display user question
        st.chat_message(
            "user"
        ).write(question)

        with st.spinner(
            "Retrieving evidence and analyzing suitability..."
        ):

            try:

                embedding_model = (
                    st.session_state.embedding_model
                )

                results = retrieve(
                    query=question,
                    index=st.session_state.index,
                    chunks=st.session_state.chunks,
                    embedding_model=embedding_model,
                    top_k=TOP_K,
                )

                # Build company profile from all company chunks
                company_chunks = [
                    chunk
                    for chunk in st.session_state.chunks
                    if chunk["document_type"]
                    == "Company Profile"
                ]

                company_profile = "\n\n".join(
                    chunk["text"]
                    for chunk in company_chunks
                )

                if not company_profile.strip():
                    st.warning(
                        "No company profile text is indexed, so an "
                        "answer would not be grounded in your company's "
                        "actual capabilities. Re-upload a text-based "
                        "company profile file and rebuild the knowledge base."
                    )
                    st.stop()

                answer, usage = analyze_with_groq(
                    question=question,
                    company_profile=company_profile,
                    retrieved_results=results,
                )

                st.session_state.messages.append(
                    {
                        "question": question,
                        "answer": answer,
                    }
                )

                st.chat_message(
                    "assistant"
                ).markdown(answer)

                is_grounded, unverified_pages = check_citation_grounding(
                    answer,
                    results,
                )

                if not is_grounded:
                    st.warning(
                        "⚠️ This answer cites page(s) "
                        f"{', '.join(str(p) for p in unverified_pages)} "
                        "that weren't in the retrieved evidence for this "
                        "question — treat that citation as unverified and "
                        "check the source document directly."
                    )

                cost = estimate_cost(usage)

                cost_col1, cost_col2, cost_col3 = st.columns(3)

                cost_col1.metric(
                    "Prompt tokens",
                    getattr(usage, "prompt_tokens", 0),
                )
                cost_col2.metric(
                    "Completion tokens",
                    getattr(usage, "completion_tokens", 0),
                )
                cost_col3.metric(
                    "Est. cost (query)",
                    f"${cost:.5f}",
                )

                st.caption(
                    "Cost is a rough estimate based on placeholder "
                    "per-token pricing — see GROQ_INPUT_COST_PER_M / "
                    "GROQ_OUTPUT_COST_PER_M in app.py."
                )

                # ----------------------------------------
                # RETRIEVED EVIDENCE
                # ----------------------------------------

                with st.expander(
                    f"🔎 Retrieved Evidence ({len(results)} chunks)"
                ):

                    for number, result in enumerate(
                        results,
                        start=1,
                    ):

                        st.markdown(
                            f"""
                            <div class="evidence-card">

                            <strong>
                            Evidence {number}
                            </strong>

                            <br>

                            <span class="small-label">
                            {result['document']}
                            • Page {result['page']}
                            • Similarity {result['similarity']:.3f}
                            </span>

                            <br><br>

                            {result['text']}

                            </div>
                            """,
                            unsafe_allow_html=True,
                        )

            except Exception as exc:

                st.error(
                    f"Analysis failed: {exc}"
                )

else:

    st.info(
        "👆 Upload both a tender document and a company profile "
        "document, then build the knowledge base."
    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "BidReady AI is an AI-assisted decision-support tool. "
    "Always verify important requirements against the original "
    "tender documentation before making a final bid decision."
)
